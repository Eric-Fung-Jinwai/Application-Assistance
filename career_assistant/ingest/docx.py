"""DOCX → raw text via python-docx (Phase 3).

Pulls both paragraph text and table-cell text, since many resume templates lay out
contact info / sections in tables.
"""

from __future__ import annotations

# Absolute import: resolves to the python-docx package, not this module.
from docx import Document


def extract_text(path: str) -> str:
    doc = Document(path)
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(part for part in parts if part.strip()).strip()
